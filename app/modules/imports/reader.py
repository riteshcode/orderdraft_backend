"""
CSV / XLSX file padhna + column mapping suggest karna.
"""
import csv
import hashlib
import io

from rapidfuzz import fuzz, process

from app.core.text import decode_bytes
from app.modules.imports.fields import FIELDS

MAX_ROWS = 50_000


class ReadError(Exception):
    pass


def read_table(content: bytes, filename: str) -> tuple[list[str], list[dict]]:
    """Returns (headers, rows). rows = [{"Header": "value"}] - sab values string."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    if name.endswith(".xls"):
        raise ReadError("The old .xls format isn't supported - in Excel, use 'Save as' to save it as .xlsx or CSV")
    return _read_csv(content)


def _read_csv(content: bytes) -> tuple[list[str], list[dict]]:
    text = decode_bytes(content)   # UTF-8 (BOM bhi), warna Windows Excel ka cp1252 - "Café" kabhi "CafÃ©" na bane
    if not text.strip():
        raise ReadError("The file is empty")

    # Delimiter detect karo (comma / semicolon / tab)
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    all_rows = [r for r in reader if any(c.strip() for c in r)]
    if not all_rows:
        raise ReadError("The file has no data")
    return _to_dicts(all_rows[0], all_rows[1:])


def _read_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ReadError(f"Couldn't open the Excel file: {e}")
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in row]
        if any(cells):
            rows.append(cells)
        if len(rows) > MAX_ROWS:
            break
    wb.close()
    if not rows:
        raise ReadError("The sheet is empty")
    return _to_dicts(rows[0], rows[1:])


def _to_dicts(header_row: list, data_rows: list[list]) -> tuple[list[str], list[dict]]:
    headers = []
    for i, h in enumerate(header_row):
        h = str(h).strip() or f"Column {i + 1}"
        while h in headers:          # duplicate header names
            h += " (2)"
        headers.append(h)
    if len(data_rows) > MAX_ROWS:
        raise ReadError(f"Max {MAX_ROWS} rows allowed")
    rows = []
    for r in data_rows:
        rows.append({h: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for i, h in enumerate(headers)})
    return headers, rows


def headers_signature(headers: list[str]) -> str:
    """Same columns wali file pehchanne ke liye (order matter nahi karta)."""
    key = "|".join(sorted(h.strip().lower() for h in headers))
    return hashlib.sha256(key.encode()).hexdigest()


def suggest_mapping(headers: list[str], kind: str) -> dict[str, str | None]:
    """
    Har CSV header ke liye best field guess karo.
      1. synonym exact match
      2. fuzzy match (>= 85)
    Ek field sirf ek column ko milta hai.
    """
    fields = FIELDS[kind]
    mapping: dict[str, str | None] = {h: None for h in headers}
    used: set[str] = set()

    # Pass 1: exact synonym
    for h in headers:
        key = h.strip().lower()
        for f in fields:
            if f["key"] not in used and (key == f["key"] or key in f["synonyms"]):
                mapping[h] = f["key"]
                used.add(f["key"])
                break

    # Pass 2: fuzzy
    choices = {}
    for f in fields:
        for syn in f["synonyms"] + [f["key"], f["label"].lower()]:
            choices[f"{f['key']}::{syn}"] = syn
    for h in headers:
        if mapping[h]:
            continue
        hit = process.extractOne(h.strip().lower(), choices, scorer=fuzz.token_sort_ratio)
        if hit and hit[1] >= 85:
            field = hit[2].split("::")[0]
            if field not in used:
                mapping[h] = field
                used.add(field)
    return mapping
